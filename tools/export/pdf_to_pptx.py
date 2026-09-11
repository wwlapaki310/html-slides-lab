#!/usr/bin/env python3
"""PDFの各ページを画像化し、1ページ=1スライドのPPTXに変換する。

    pip install pymupdf python-pptx
    python tools/export/pdf_to_pptx.py exports/2026/sample-slide.pdf
    python tools/export/pdf_to_pptx.py exports            # フォルダ配下の全PDFを一括変換

注意: 画像を貼り付けたPPTXになるため、PowerPoint上でテキストを直接編集することはできない。
「配布・共有」用途なら十分だが、「PowerPointで編集」が必要な場合は手動移植が必要。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import fitz  # PyMuPDF
    from pptx import Presentation
    from pptx.util import Emu
except ImportError as e:  # pragma: no cover
    sys.exit(f"依存パッケージが不足しています ({e.name})。`pip install pymupdf python-pptx` を実行してください。")

EMU_PER_INCH = 914400


def convert(pdf_path: Path, out_path: Path, dpi: int = 150) -> None:
    doc = fitz.open(pdf_path)
    if doc.page_count == 0:
        raise ValueError(f"ページがありません: {pdf_path}")

    # 1ページ目のサイズ（pt）をそのままスライドサイズにする → 余白なしで貼り付く
    first = doc.load_page(0)
    prs = Presentation()
    prs.slide_width = Emu(int(first.rect.width / 72 * EMU_PER_INCH))
    prs.slide_height = Emu(int(first.rect.height / 72 * EMU_PER_INCH))
    blank_layout = prs.slide_layouts[6]  # 完全な白紙レイアウト

    tmp_dir = out_path.parent / f".{out_path.stem}_pages"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    try:
        for i in range(doc.page_count):
            pix = doc.load_page(i).get_pixmap(dpi=dpi)
            img_path = tmp_dir / f"page-{i + 1:04d}.png"
            pix.save(img_path)
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)
        print(f"{pdf_path} -> {out_path} ({doc.page_count}ページ)")
    finally:
        doc.close()
        for f in tmp_dir.glob("page-*.png"):
            f.unlink()
        tmp_dir.rmdir()


def main() -> int:
    ap = argparse.ArgumentParser(description="PDF を 1ページ1画像の PPTX に変換する")
    ap.add_argument("target", type=Path, help="PDFファイル、またはPDFを含むフォルダ")
    ap.add_argument("--dpi", type=int, default=150, help="ページ画像の解像度（既定: 150）")
    args = ap.parse_args()

    pdfs = sorted(args.target.rglob("*.pdf")) if args.target.is_dir() else [args.target]
    if not pdfs:
        print(f"PDFが見つかりません: {args.target}", file=sys.stderr)
        return 1

    for pdf in pdfs:
        convert(pdf, pdf.with_suffix(".pptx"), dpi=args.dpi)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
